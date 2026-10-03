"""Pembuat laporan bulanan (Excel & PDF) - Mahkota Graha KPR."""
from io import BytesIO
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

BLUE = "1E3A8A"
LIGHT = "DBEAFE"

LEGAL_LABELS = {
    "status_sertifikat": "Sertifikat",
    "status_imb_pbg": "IMB/PBG",
    "status_pbb": "PBB",
    "status_ssp_pph": "SSP/PPh",
    "status_bphtb": "BPHTB",
}


def _period_label(ctx: dict) -> str:
    m = ctx.get("month")
    if not m:
        return "Semua Periode"
    try:
        d = datetime.strptime(m, "%Y-%m")
        bulan = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli",
                 "Agustus", "September", "Oktober", "November", "Desember"]
        return f"{bulan[d.month - 1]} {d.year}"
    except Exception:
        return m


def _title(ctx: dict) -> str:
    t = f"Laporan KPR {_period_label(ctx)}"
    if ctx.get("marketing"):
        t += f" - Marketing {ctx['marketing']}"
    return t


def _kpi_rows(ctx: dict):
    k = ctx["kpi"]
    u = ctx["unit_summary"]
    sections = ctx.get("sections", ("kpr", "unit", "legal"))
    rows = []
    if "kpr" in sections:
        rows += [
            ("Total Berkas", k["total"]), ("Proses", k["proses"]), ("SP3K", k["sp3k"]),
            ("Done (Akad)", k["done"]), ("Diputihkan", k["diputihkan"]),
            ("Mendekati Jatuh Tempo", k["nearing_count"]),
            ("% SP3K", f"{k['pct_sp3k']}%"), ("% Done", f"{k['pct_done']}%"),
        ]
    if "unit" in sections:
        rows += [
            ("Total Unit", u["total"]), ("Unit Proses", u["proses"]), ("Unit Terlambat", u["terlambat"]),
            ("Unit Done", u["done"]), ("Rata-rata Progres Bangunan", f"{u['avg_progress']}%"),
            ("Unit Siap Dipasarkan", u["siap_dipasarkan"]),
        ]
    if "legal" in sections:
        for f, label in LEGAL_LABELS.items():
            counts = ctx.get("legal_summary", {}).get(f, {})
            rows.append((f"Legalitas {label}", ", ".join(f"{s}: {n}" for s, n in counts.items()) or "-"))
    return rows


def _kpr_table(ctx: dict):
    head = ["No", "Nama Konsumen", "Blok", "Marketing", "Bank", "Tgl Booking",
            "Tahap", "Status", "Tgl SP3K", "Tgl Akad", "Sisa Hari", "Keterangan"]
    rows = []
    for i, r in enumerate(sorted(ctx["kpr_rows"], key=lambda x: x.get("tanggal_booking") or ""), 1):
        sisa = r.get("days_to_pemutihan")
        rows.append([
            i, r.get("nama_konsumen", ""), r.get("blok_kavling", ""), r.get("marketing", ""),
            r.get("bank_pemroses", ""), r.get("tanggal_booking", ""), r.get("tahap_saat_ini", ""),
            r.get("status", ""), r.get("tanggal_sp3k") or "-", r.get("tanggal_akad") or "-",
            sisa if (r.get("status") == "PROSES" and sisa is not None) else "-",
            r.get("keterangan", "") or "",
        ])
    return head, rows


def _marketing_table(ctx: dict):
    head = ["Marketing", "Total", "Proses", "SP3K + Done", "Diputihkan", "Rata² Progres Unit"]
    rows = [[m["marketing"], m["total"], m["proses"], m["sp3k_done"], m["diputihkan"], f"{m['avg_progress']}%"]
            for m in sorted(ctx["by_marketing"], key=lambda x: -x["total"])]
    return head, rows


def _bank_table(ctx: dict):
    head = ["Bank", "Total", "Proses", "SP3K + Done"]
    rows = [[b["bank"], b["total"], b["proses"], b["sp3k_done"]]
            for b in sorted(ctx["by_bank"], key=lambda x: -x["total"])]
    return head, rows


def _unit_table(ctx: dict):
    head = ["Blok", "Kontraktor", "Tahap", "Progres", "Status", "Target Selesai", "Kendala"]
    rows = [[u.get("blok_kavling", ""), u.get("nama_kontraktor", ""), u.get("tahap_konstruksi", ""),
             f"{u.get('persen_progres', 0)}%", u.get("status_bangunan", ""),
             u.get("tanggal_target_selesai") or "-", u.get("kendala_catatan") or ""]
            for u in sorted(ctx["units"], key=lambda x: x.get("blok_kavling", ""))]
    return head, rows


def _legal_table(ctx: dict):
    head = ["Blok", "Sertifikat", "IMB/PBG", "PBB", "SSP/PPh", "BPHTB", "Keterangan"]
    rows = [[l.get("blok_kavling", ""), l.get("status_sertifikat") or "-", l.get("status_imb_pbg") or "-",
             l.get("status_pbb") or "-", l.get("status_ssp_pph") or "-", l.get("status_bphtb") or "-",
             l.get("keterangan") or ""]
            for l in sorted(ctx["legality"], key=lambda x: x.get("blok_kavling", ""))]
    return head, rows


# ---------------- Excel ----------------
def _write_sheet(ws, title: str, head, rows, start_row: int = 1):
    ws.cell(row=start_row, column=1, value=title).font = Font(bold=True, size=13, color=BLUE)
    r = start_row + 1
    for c, h in enumerate(head, 1):
        cell = ws.cell(row=r, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        r += 1
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=v)
    if not rows:
        r += 1
        ws.cell(row=r, column=1, value="Tidak ada data")
    for c in range(1, len(head) + 1):
        width = max([len(str(head[c - 1]))] + [len(str(row[c - 1])) for row in rows] + [8])
        ws.column_dimensions[get_column_letter(c)].width = min(width + 2, 45)
    return r + 2


def build_xlsx(ctx: dict) -> bytes:
    sections = ctx.get("sections", ("kpr", "unit", "legal"))
    wb = Workbook()
    ws = wb.active
    ws.title = "Ringkasan"
    ws["A1"] = ctx["company_name"]
    ws["A1"].font = Font(bold=True, size=14, color=BLUE)
    ws["A2"] = ctx["project_name"]
    ws["A3"] = _title(ctx)
    ws["A3"].font = Font(bold=True, size=12)
    ws["A4"] = f"Dibuat: {ctx['generated_at'][:19].replace('T', ' ')} UTC · oleh {ctx.get('generated_by', '-')}"
    r = 6
    for label, val in _kpi_rows(ctx):
        ws.cell(row=r, column=1, value=label).fill = PatternFill("solid", fgColor=LIGHT)
        ws.cell(row=r, column=2, value=val).font = Font(bold=True)
        r += 1
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 14
    r += 1
    if "kpr" in sections:
        r = _write_sheet(ws, "Rekap per Marketing", *_marketing_table(ctx), start_row=r)
        _write_sheet(ws, "Rekap per Bank", *_bank_table(ctx), start_row=r)
        _write_sheet(wb.create_sheet("Berkas KPR"), _title(ctx), *_kpr_table(ctx))
    if "unit" in sections:
        _write_sheet(wb.create_sheet("Progres Bangunan"), "Progres Bangunan per Unit", *_unit_table(ctx))
    if "legal" in sections:
        _write_sheet(wb.create_sheet("Legalitas"), "Legalitas per Unit", *_legal_table(ctx))

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------- PDF ----------------
def _pdf_table(head, rows, col_widths=None, font_size=7):
    body = rows if rows else [["Tidak ada data"] + [""] * (len(head) - 1)]
    cell = ParagraphStyle("cell", fontSize=font_size, leading=font_size + 2)
    hstyle = ParagraphStyle("head", fontSize=font_size, leading=font_size + 2,
                            textColor=colors.white, fontName="Helvetica-Bold")
    data = [[Paragraph(str(v), hstyle) for v in head]]
    data += [[Paragraph(str(v), cell) for v in row] for row in body]
    t = Table(data, colWidths=col_widths, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + BLUE)),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EFF6FF")]),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
    ]))
    return t


def build_pdf(ctx: dict) -> bytes:
    sections = ctx.get("sections", ("kpr", "unit", "legal"))
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm,
                            title=_title(ctx), author=ctx["company_name"])
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Title"], fontSize=16, textColor=colors.HexColor("#" + BLUE),
                        alignment=0, spaceAfter=2)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontSize=11, textColor=colors.HexColor("#" + BLUE),
                        spaceBefore=8, spaceAfter=4)
    small = ParagraphStyle("small", parent=ss["Normal"], fontSize=8, textColor=colors.HexColor("#64748B"))
    W = doc.width

    story = [
        Paragraph(ctx["company_name"], h1),
        Paragraph(f"{ctx['project_name']} · {_title(ctx)}", ss["Normal"]),
        Paragraph(f"Dibuat: {ctx['generated_at'][:19].replace('T', ' ')} UTC · oleh {ctx.get('generated_by', '-')}", small),
        Spacer(1, 6),
        Paragraph("Ringkasan", h2),
    ]
    kpis = _kpi_rows(ctx)
    half = (len(kpis) + 1) // 2
    kpi_rows = []
    for i in range(half):
        left = kpis[i]
        right = kpis[i + half] if i + half < len(kpis) else ("", "")
        kpi_rows.append([left[0], left[1], right[0], right[1]])
    story.append(_pdf_table(["Indikator", "Nilai", "Indikator", "Nilai"], kpi_rows,
                            col_widths=[W * 0.25, W * 0.25, W * 0.25, W * 0.25], font_size=8))

    if "kpr" in sections:
        story.append(Paragraph("Rekap per Marketing", h2))
        story.append(_pdf_table(*_marketing_table(ctx), font_size=8))
        story.append(Paragraph("Rekap per Bank", h2))
        story.append(_pdf_table(*_bank_table(ctx), font_size=8))
        story.append(Paragraph("Daftar Berkas KPR", h2))
        head, rows = _kpr_table(ctx)
        widths = [W * w for w in (0.03, 0.14, 0.06, 0.09, 0.07, 0.08, 0.09, 0.08, 0.08, 0.08, 0.05, 0.15)]
        story.append(_pdf_table(head, rows, col_widths=widths))

    if "unit" in sections:
        story.append(Paragraph("Progres Bangunan", h2))
        head, rows = _unit_table(ctx)
        story.append(_pdf_table(head, rows, col_widths=[W * w for w in (0.08, 0.17, 0.15, 0.08, 0.12, 0.12, 0.28)]))

    if "legal" in sections:
        story.append(Paragraph("Legalitas per Unit", h2))
        head, rows = _legal_table(ctx)
        story.append(_pdf_table(head, rows, col_widths=[W * w for w in (0.1, 0.13, 0.13, 0.13, 0.13, 0.13, 0.25)]))

    doc.build(story)
    return buf.getvalue()
